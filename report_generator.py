# 声明：本代码仅供学习和研究目的使用。使用者应遵守以下原则：
# 1. 不得用于任何商业用途。
# 2. 使用时应遵守目标平台的使用条款和robots.txt规则。
# 3. 不得进行大规模爬取或对平台造成运营干扰。
# 4. 应合理控制请求频率，避免给目标平台带来不必要的负担。
# 5. 不得用于任何非法或不当的用途。

"""
报告生成器
支持 JSON、CSV、Excel、HTML 格式的报告导出
"""

import json
import csv
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional
import io


class ReportGenerator:
    def __init__(self, reports_dir: str = "./reports"):
        self.reports_dir = Path(reports_dir)
        self.reports_dir.mkdir(parents=True, exist_ok=True)
    
    def generate_json_report(self, data: List[Dict[str, Any]], job_id: str) -> Path:
        """生成 JSON 报告"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{job_id}_{timestamp}.json"
        filepath = self.reports_dir / filename
        
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        return filepath
    
    def generate_csv_report(self, data: List[Dict[str, Any]], job_id: str) -> Path:
        """生成 CSV 报告"""
        if not data:
            raise ValueError("没有数据可导出")
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{job_id}_{timestamp}.csv"
        filepath = self.reports_dir / filename
        
        fieldnames = list(data[0].keys())
        
        with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data)
        
        return filepath
    
    def generate_excel_report(self, data: List[Dict[str, Any]], job_id: str, 
                             sheets: Optional[Dict[str, List[Dict]]] = None) -> Path:
        """生成 Excel 报告（支持多个工作表）"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            raise ImportError("需要安装 openpyxl: pip install openpyxl")
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{job_id}_{timestamp}.xlsx"
        filepath = self.reports_dir / filename
        
        wb = Workbook()
        wb.remove(wb.active)
        
        # 如果提供了多个工作表
        if sheets:
            for sheet_name, sheet_data in sheets.items():
                if not sheet_data:
                    continue
                self._add_excel_sheet(wb, sheet_name, sheet_data)
        else:
            self._add_excel_sheet(wb, "数据", data)
        
        wb.save(filepath)
        return filepath
    
    def _add_excel_sheet(self, workbook, sheet_name: str, data: List[Dict[str, Any]]):
        """添加 Excel 工作表"""
        from openpyxl.styles import Font, Alignment, PatternFill
        
        ws = workbook.create_sheet(title=sheet_name)
        
        if not data:
            return
        
        # 表头
        headers = list(data[0].keys())
        ws.append(headers)
        
        # 设置表头样式
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 数据行
        for item in data:
            row = [item.get(h, "") for h in headers]
            ws.append(row)
        
        # 自动调整列宽
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def generate_html_report(self, data: List[Dict[str, Any]], job_id: str,
                            template: Optional[str] = None) -> Path:
        """生成 HTML 报告"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{job_id}_{timestamp}.html"
        filepath = self.reports_dir / filename
        
        if template:
            html_content = template
        else:
            html_content = self._generate_default_html(data, job_id)
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        return filepath
    
    def _generate_default_html(self, data: List[Dict[str, Any]], job_id: str) -> str:
        """生成默认 HTML 模板"""
        if not data:
            return "<html><body><h1>无数据</h1></body></html>"
        
        headers = list(data[0].keys())
        
        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>MediaCrawler 报告 - {job_id}</title>
    <style>
        body {{
            font-family: 'Arial', sans-serif;
            margin: 20px;
            background: #f5f5f5;
        }}
        .container {{
            max-width: 1400px;
            margin: 0 auto;
            background: white;
            padding: 30px;
            border-radius: 8px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }}
        h1 {{
            color: #333;
            border-bottom: 3px solid #4472C4;
            padding-bottom: 10px;
        }}
        .meta {{
            color: #666;
            margin-bottom: 20px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
        }}
        th {{
            background: #4472C4;
            color: white;
            padding: 12px;
            text-align: left;
            font-weight: bold;
        }}
        td {{
            padding: 10px 12px;
            border-bottom: 1px solid #ddd;
        }}
        tr:hover {{
            background: #f5f5f5;
        }}
        .stats {{
            display: flex;
            gap: 20px;
            margin: 20px 0;
        }}
        .stat-card {{
            flex: 1;
            padding: 20px;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            border-radius: 8px;
        }}
        .stat-value {{
            font-size: 32px;
            font-weight: bold;
        }}
        .stat-label {{
            font-size: 14px;
            opacity: 0.9;
        }}
    </style>
</head>
<body>
    <div class="container">
        <h1>📊 MediaCrawler 数据报告</h1>
        <div class="meta">
            <strong>任务ID:</strong> {job_id} | 
            <strong>生成时间:</strong> {datetime.now().strftime("%Y-%m-%d %H:%M:%S")} | 
            <strong>数据条数:</strong> {len(data)}
        </div>
        
        <div class="stats">
            <div class="stat-card">
                <div class="stat-value">{len(data)}</div>
                <div class="stat-label">总记录数</div>
            </div>
            <div class="stat-card" style="background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);">
                <div class="stat-value">{len(headers)}</div>
                <div class="stat-label">字段数量</div>
            </div>
        </div>
        
        <table>
            <thead>
                <tr>
                    {"".join(f"<th>{h}</th>" for h in headers)}
                </tr>
            </thead>
            <tbody>
                {"".join("<tr>" + "".join(f"<td>{item.get(h, '')}</td>" for h in headers) + "</tr>" for item in data)}
            </tbody>
        </table>
    </div>
</body>
</html>"""
        return html
    
    def generate_summary_report(self, job_data: Dict[str, Any]) -> Dict[str, Any]:
        """生成任务摘要报告"""
        return {
            "job_id": job_data.get("job_id"),
            "platform": job_data.get("config", {}).get("platform"),
            "keywords": job_data.get("config", {}).get("keywords"),
            "status": job_data.get("status"),
            "start_time": job_data.get("start_time"),
            "end_time": job_data.get("end_time"),
            "notes_crawled": job_data.get("progress", {}).get("notes_crawled", 0),
            "comments_crawled": job_data.get("progress", {}).get("comments_crawled", 0),
            "error": job_data.get("error"),
            "generated_at": datetime.now().isoformat()
        }


# 全局报告生成器
report_generator = ReportGenerator()
